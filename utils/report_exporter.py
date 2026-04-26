"""
Generates a professionally formatted .docx fault analysis report.
"""
from __future__ import annotations

import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from data.simulator import SensorSnapshot, METRICS


def generate_report(analysis: dict) -> bytes:
    """
    Build the .docx report from the completed agent analysis dict.
    Returns bytes suitable for st.download_button.
    """
    doc = Document()
    _set_page_margins(doc)

    scenario: dict         = analysis["scenario"]
    snapshot: SensorSnapshot = analysis["snapshot"]
    triage_text: str        = analysis.get("triage", "")
    primary_name: str       = analysis.get("expert_primary_name", "")
    primary_text: str       = analysis.get("expert_primary_text", "")
    secondary_name: str     = analysis.get("expert_secondary_name", "")
    secondary_text: str     = analysis.get("expert_secondary_text", "")
    faulty_steps: list      = analysis.get("faulty_steps", [])

    timestamp = datetime.now().strftime("%B %d, %Y  %H:%M:%S")

    # ── Title block ──────────────────────────────────────────────────────────
    title = doc.add_heading("Fault Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_color(title.runs[0], RGBColor(0x1A, 0x3A, 0x5C))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Near-Closed-Loop PV Manufacturing Control System").bold = True
    doc.add_paragraph().add_run(f"Generated: {timestamp}").italic = True
    doc.add_paragraph().add_run(
        f"Powered by Claude {_get_model_short()} + Semantic Kernel"
    ).italic = True

    doc.add_paragraph()

    # ── Executive Summary ────────────────────────────────────────────────────
    doc.add_heading("Executive Summary", level=1)
    summary = (
        f"A '{scenario['display_name']}' fault was detected at "
        f"Stage {scenario['stage']} ({scenario['stage_name']}) of the PV manufacturing line "
        f"on {timestamp}. "
        f"The centralized reasoning agent was automatically triggered and performed a multi-step "
        f"root cause analysis, consulting specialist expert agents to diagnose the fault. "
        f"Severity: {scenario.get('severity', 'Unknown')}."
    )
    doc.add_paragraph(summary)

    # ── Detected Anomaly ─────────────────────────────────────────────────────
    doc.add_heading("Detected Anomaly", level=1)
    doc.add_paragraph(f"Fault scenario: {scenario['display_name']}")
    doc.add_paragraph(
        f"Plain-language description: {scenario.get('plain_english', scenario.get('description', ''))}"
    )

    if faulty_steps:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, label in enumerate(["Step", "Parameter", "Unit", "Observed Value", "Threshold (Fault)"]):
            hdr[i].text = label
            hdr[i].paragraphs[0].runs[0].bold = True

        for f in faulty_steps:
            row = table.add_row().cells
            row[0].text = f["step"]
            row[1].text = f["name"]
            row[2].text = f["unit"]
            row[3].text = str(f["value"])
            row[4].text = f"{f['direction']} {f['threshold']}"

    doc.add_paragraph()

    # ── Root Cause Analysis (Triage) ─────────────────────────────────────────
    doc.add_heading("Root Cause Analysis — Central Reasoning Agent", level=1)
    if triage_text:
        for para in triage_text.split("\n\n"):
            para = para.strip()
            if para.startswith("## "):
                doc.add_heading(para[3:], level=2)
            elif para.startswith("# "):
                doc.add_heading(para[2:], level=2)
            elif para:
                doc.add_paragraph(para)
    else:
        doc.add_paragraph("Triage analysis not available.")

    # ── Primary Expert Findings ──────────────────────────────────────────────
    doc.add_heading(f"Expert Findings — {primary_name}", level=1)
    _add_loop_badge(doc, analysis, "primary")
    if primary_text:
        for para in primary_text.split("\n\n"):
            para = para.strip()
            if para.startswith("## "):
                doc.add_heading(para[3:], level=2)
            elif para:
                doc.add_paragraph(para)
    else:
        doc.add_paragraph("Expert analysis not available.")

    # ── Secondary Expert Findings ────────────────────────────────────────────
    if secondary_text and secondary_name:
        doc.add_heading(f"Cross-Check Findings — {secondary_name}", level=1)
        _add_loop_badge(doc, analysis, "secondary")
        for para in secondary_text.split("\n\n"):
            para = para.strip()
            if para.startswith("## "):
                doc.add_heading(para[3:], level=2)
            elif para:
                doc.add_paragraph(para)

    # ── Full Sensor Snapshot ─────────────────────────────────────────────────
    doc.add_heading("Sensor Data Snapshot at Time of Fault", level=1)
    snap_table = doc.add_table(rows=1, cols=5)
    snap_table.style = "Table Grid"
    hdr = snap_table.rows[0].cells
    for i, label in enumerate(["Step", "Parameter", "Value", "Unit", "Status"]):
        hdr[i].text = label
        hdr[i].paragraphs[0].runs[0].bold = True

    for step_key, val in snapshot.metrics.items():
        meta = METRICS[step_key]
        from data.simulator import _check_status
        status = _check_status(step_key, val)
        row = snap_table.add_row().cells
        row[0].text = step_key
        row[1].text = meta["name"]
        row[2].text = str(round(val, 4))
        row[3].text = meta["unit"]
        row[4].text = status

    # ── Footer ───────────────────────────────────────────────────────────────
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run(
        "Generated by Near-Closed-Loop PV Manufacturing Control System  |  "
        f"Model: Claude {_get_model_short()}  |  Orchestrator: Semantic Kernel"
    )
    footer_run.italic = True
    footer_run.font.size = Pt(8)
    _set_run_color(footer_run, RGBColor(0x88, 0x99, 0xAA))

    # Serialize to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _set_run_color(run, color: RGBColor):
    run.font.color.rgb = color


def _set_page_margins(doc: Document):
    from docx.shared import Inches
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.2)
        section.right_margin  = Inches(1.2)


def _add_loop_badge(doc: Document, analysis: dict, which: str):
    from agents.triage_agent import EXPERT_MAP
    key = analysis["scenario"].get(f"{'primary' if which == 'primary' else 'secondary'}_expert", "")
    cls = EXPERT_MAP.get(key)
    if cls:
        p = doc.add_paragraph()
        p.add_run(f"Control Loop: {cls.loop_steps}").italic = True
        p.add_run(f"   |   ").font.size = Pt(8)


def _get_model_short() -> str:
    import os
    from dotenv import load_dotenv
    load_dotenv()
    return os.getenv("MODEL_ID", "claude-sonnet-4-6")
